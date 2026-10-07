"""Unified loader tests — including the two ingestion bugs found in the audit.

The audit found that (a) XLSX formula detection never fired because the workbook
was opened with ``data_only=True``, and (b) every negative number was treated as a
formula-injection payload because numeric cells were stringified and then
prefix-tested, causing legitimate workbooks to be rejected wholesale.

Both regressions are guarded here, along with the structural intelligence that was
missing entirely (merged cells, hidden rows/columns) and the macro-extension hole.
"""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from app.services.orbit.contracts import ErrorCategory, IngestionError
from app.services.orbit.ingestion.loaders import load_artifact


def _xlsx(tmp_path, rows, name="book.xlsx", sheet="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path.read_bytes()


# ── CSV ──────────────────────────────────────────────────────────────────────

class TestCsv:
    def test_reads_headers_and_data(self, tmp_path):
        path = tmp_path / "sales.csv"
        path.write_text(
            "Item,Qty,Price,Date\n"
            "Al Noor Milk,12,4.50,2026-09-01\n"
            "Pepsi 330ml,30,2.00,2026-09-02\n",
            encoding="utf-8",
        )
        art = load_artifact(path.read_bytes(), "sales.csv")
        sheet = art.primary_sheet
        assert art.kind == "csv"
        assert sheet.headers == ["Item", "Qty", "Price", "Date"]
        assert sheet.header_index == 0
        assert len(sheet.data_rows) == 2

    def test_bom_and_arabic_headers(self, tmp_path):
        path = tmp_path / "arabic.csv"
        path.write_text("الصنف,الكمية,السعر\nحليب النور,12,4.50\n", encoding="utf-8-sig")
        art = load_artifact(path.read_bytes(), "arabic.csv")
        assert art.primary_sheet.headers[0] == "الصنف"

    def test_semicolon_delimiter(self, tmp_path):
        path = tmp_path / "semi.csv"
        path.write_text("Item;Qty\nA;1\nB;2\n", encoding="utf-8")
        art = load_artifact(path.read_bytes(), "semi.csv")
        assert art.primary_sheet.headers == ["Item", "Qty"]

    def test_currency_hint_is_a_hint_not_a_conversion(self, tmp_path):
        path = tmp_path / "cur.csv"
        path.write_text("Item,Total\nA,100 USD\nB,250 SAR\n", encoding="utf-8")
        sheet = load_artifact(path.read_bytes(), "cur.csv").primary_sheet
        # A hint only; per-cell currency is the currency normalizer's job.
        assert sheet.metadata.currency_hint in {"USD", "SAR"}


# ── Audit bug 2: negative numbers are not formulas ──────────────────────────

class TestNegativeNumbersAreNotFormulas:
    def test_negative_numbers_load_without_rejection(self, tmp_path):
        path = tmp_path / "ledger.csv"
        path.write_text("Item,Delta\nRefund,-5\nSale,12\nAdjustment,-30\n", encoding="utf-8")
        sheet = load_artifact(path.read_bytes(), "ledger.csv").primary_sheet
        values = [r[1] for r in sheet.data_rows]
        assert "-5" in values and "-30" in values

    def test_negative_numbers_in_xlsx_survive(self, tmp_path):
        content = _xlsx(tmp_path, [["Item", "Qty"], ["Credit", -5], ["Sale", 12]])
        sheet = load_artifact(content, "book.xlsx").primary_sheet
        records = sheet.to_records()
        assert any(r.get("Item") == "Credit" and str(r.get("Qty")) == "-5" for r in records)


# ── Audit bug 1: XLSX formulas are detected ─────────────────────────────────

class TestFormulaDetection:
    def test_formula_cells_are_detected(self, tmp_path):
        content = _xlsx(tmp_path, [["Item", "Qty"], ["Milk", "=SUM(1,2,3)"], ["Juice", 10]])
        meta = load_artifact(content, "book.xlsx").primary_sheet.metadata
        assert meta.formula_cells, "formula cells were not detected"
        assert any("formula" in w for w in meta.warnings)

    def test_ordinary_formulas_are_not_treated_as_external(self, tmp_path):
        content = _xlsx(
            tmp_path,
            [["A", "B", "C"], [10, 20, "=SUM(A2:B2)"], [1, 2, "=IF(A3>0,C3,0)"]],
        )
        sheet = load_artifact(content, "book.xlsx").primary_sheet
        assert not sheet.metadata.external_links

    def test_formula_without_cached_value_is_omitted_not_zero(self, tmp_path):
        """A missing value must not become 0."""
        content = _xlsx(tmp_path, [["Item", "Qty"], ["Milk", "=SUM(1,2)"]])
        records = load_artifact(content, "book.xlsx").primary_sheet.to_records()
        milk = next(r for r in records if r.get("Item") == "Milk")
        assert "Qty" not in milk, "a missing formula value was coerced to zero"


# ── Structural intelligence that was entirely missing ───────────────────────

class TestStructuralCapture:
    def _structural(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Stock"
        ws.append(["Item", "Qty", "Note"])
        ws.append(["A", 1, ""])
        ws.append(["B", 2, ""])
        ws.merge_cells("A3:C3")
        ws.cell(row=3, column=1, value="merged note")
        ws.row_dimensions[2].hidden = True
        ws.column_dimensions["B"].hidden = True
        path = tmp_path / "structure.xlsx"
        wb.save(path)
        return path.read_bytes()

    def test_merged_cells_captured(self, tmp_path):
        meta = load_artifact(self._structural(tmp_path), "s.xlsx").primary_sheet.metadata
        assert meta.merged_cell_ranges

    def test_hidden_rows_and_columns_captured(self, tmp_path):
        meta = load_artifact(self._structural(tmp_path), "s.xlsx").primary_sheet.metadata
        assert meta.hidden_rows
        assert meta.hidden_columns

    def test_hidden_worksheet_refused(self, tmp_path):
        wb = Workbook()
        wb.active.title = "visible"
        hidden = wb.create_sheet("hidden_sheet")
        hidden["A1"] = "=HYPERLINK(\"http://evil.test\",\"x\")"
        wb["hidden_sheet"].sheet_state = "hidden"
        path = tmp_path / "hidden.xlsx"
        wb.save(path)
        with pytest.raises(IngestionError) as exc:
            load_artifact(path.read_bytes(), "hidden.xlsx")
        assert exc.value.category is ErrorCategory.SECURITY_ERROR


# ── Security ────────────────────────────────────────────────────────────────

class TestSecurity:
    @pytest.mark.parametrize("ext", ["xlsm", "xlsb", "xlam"])
    def test_macro_extensions_refused(self, ext):
        with pytest.raises(IngestionError) as exc:
            load_artifact(b"PK\x03\x04junk", f"book.{ext}")
        assert exc.value.category is ErrorCategory.SECURITY_ERROR

    def test_external_hyperlink_formula_refused(self, tmp_path):
        content = _xlsx(tmp_path, [["Item", "Link"], ["A", '=HYPERLINK("http://evil.test","x")']])
        with pytest.raises(IngestionError) as exc:
            load_artifact(content, "book.xlsx")
        assert exc.value.category is ErrorCategory.SECURITY_ERROR

    def test_workbook_link_formula_refused(self, tmp_path):
        content = _xlsx(tmp_path, [["Item", "Ref"], ["A", "='[other.xlsx]Sheet1'!A1"]])
        with pytest.raises(IngestionError) as exc:
            load_artifact(content, "book.xlsx")
        assert exc.value.category is ErrorCategory.SECURITY_ERROR

    def test_dangerous_csv_content_refused(self, tmp_path):
        path = tmp_path / "x.csv"
        path.write_text("Item,Note\nA,<script>alert(1)</script>\n", encoding="utf-8")
        with pytest.raises(IngestionError) as exc:
            load_artifact(path.read_bytes(), "x.csv")
        assert exc.value.category is ErrorCategory.SECURITY_ERROR

    def test_unsupported_format_is_categorised(self):
        with pytest.raises(IngestionError) as exc:
            load_artifact(b"%PDF-1.4", "invoice.pdf")
        assert exc.value.category is ErrorCategory.UNSUPPORTED_FORMAT

    def test_row_limit_enforced(self, monkeypatch):
        monkeypatch.setattr("app.services.orbit.ingestion.loaders.MAX_ROWS", 3)
        body = "Item,Qty\n" + "".join(f"A{i},{i}\n" for i in range(20))
        with pytest.raises(IngestionError) as exc:
            load_artifact(body.encode(), "big.csv")
        assert exc.value.category is ErrorCategory.SCHEMA_ERROR

    def test_malformed_json_is_parse_error(self):
        with pytest.raises(IngestionError) as exc:
            load_artifact(b"{not json", "broken.json")
        assert exc.value.category is ErrorCategory.PARSE_ERROR


# ── JSON ────────────────────────────────────────────────────────────────────

class TestJson:
    def test_array_of_objects(self, tmp_path):
        path = tmp_path / "d.json"
        path.write_text('[{"item":"Milk","qty":5},{"item":"Juice","qty":2}]', encoding="utf-8")
        sheet = load_artifact(path.read_bytes(), "d.json").primary_sheet
        assert sheet.headers == ["item", "qty"]
        assert len(sheet.data_rows) == 2

    def test_object_of_arrays_becomes_sheets(self, tmp_path):
        path = tmp_path / "d.json"
        path.write_text('{"sales":[{"a":1}],"stock":[{"b":2}]}', encoding="utf-8")
        art = load_artifact(path.read_bytes(), "d.json")
        assert {s.name for s in art.sheets} == {"sales", "stock"}


# ── Idempotent shape ────────────────────────────────────────────────────────

class TestDeterministicLoading:
    def test_same_bytes_produce_identical_structure(self, tmp_path):
        path = tmp_path / "s.csv"
        path.write_text("Item,Qty\nA,1\nB,2\n", encoding="utf-8")
        content = path.read_bytes()
        a = load_artifact(content, "s.csv").primary_sheet
        b = load_artifact(content, "s.csv").primary_sheet
        assert a.headers == b.headers
        assert a.data_rows == b.data_rows

    def test_duplicate_headers_are_made_unique(self, tmp_path):
        path = tmp_path / "d.csv"
        path.write_text("Item,Item,Item\nA,1,2\n", encoding="utf-8")
        sheet = load_artifact(path.read_bytes(), "d.csv").primary_sheet
        assert len(set(sheet.headers)) == len(sheet.headers)