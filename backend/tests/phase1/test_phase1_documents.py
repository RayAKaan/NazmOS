"""Document extraction tests (spec §11).

PDF/DOCX/text extraction did not exist before Phase 1. These tests pin the
behaviour that matters:

* every extracted item carries a source locator (traceability);
* only monetary fields carry an amount (an invoice number is not 2,026 SAR);
* document titles are not mistaken for field values ("TAX INVOICE" is not a tax
  amount, "Supplier Quotation" is not a supplier name);
* unrecognised lines are preserved rather than discarded;
* a document with no text layer is reported honestly instead of guessed at.
"""
from __future__ import annotations

import io

import pytest
from docx import Document

from app.services.orbit.contracts import ErrorCategory, ExtractionMethod, IngestionError
from app.services.orbit.ingestion.documents import extract_document


def _docx_bytes(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _minimal_pdf(text: str = "Invoice INV-X") -> bytes:
    """A tiny hand-built single-page PDF (no external dependency needed)."""
    content = b"BT /F1 12 Tf 60 780 Td (" + text.encode("latin-1", "replace") + b") Tj ET"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += (
        b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\n"
        b"startxref\n" + str(xref).encode() + b"\n%%EOF\n"
    )
    return bytes(out)


def _quote_doc() -> bytes:
    def build(doc: Document) -> None:
        doc.add_heading("Supplier Quotation", level=1)
        doc.add_paragraph("Quote No: Q-2026-0042")
        doc.add_paragraph("Supplier: Al Noor Trading Est.")
        doc.add_paragraph("Date: 2026-09-30")
        doc.add_paragraph("Delivery terms are negotiable.")
        t = doc.add_table(rows=3, cols=4)
        for i, row in enumerate(
            [
                ["Item", "Qty", "Unit Price", "Total"],
                ["Item A", "12", "14.00", "168.00"],
                ["Item B", "5", "9.50", "47.50"],
            ]
        ):
            for j, v in enumerate(row):
                t.cell(i, j).text = v

    return _docx_bytes(build)


# ── DOCX ────────────────────────────────────────────────────────────────────

class TestDocx:
    def test_extracts_labelled_fields(self):
        doc = extract_document(_quote_doc(), "quote.docx")
        assert doc.doc_type == "docx"
        names = {f.field_name for f in doc.fields}
        assert "quote_number" in names
        assert "supplier_name" in names

    def test_extracts_tables(self):
        doc = extract_document(_quote_doc(), "quote.docx")
        assert len(doc.tables) == 1
        assert doc.tables[0].row_count == 3
        assert doc.tables[0].column_count == 4

    def test_every_field_has_a_locator(self):
        doc = extract_document(_quote_doc(), "quote.docx")
        for f in doc.fields:
            assert f.locator.row is not None or f.locator.page is not None, (
                f"{f.field_name} has no source locator"
            )

    def test_document_title_is_not_a_field_value(self):
        """"Supplier Quotation" must not yield supplier_name='Quotation'."""
        doc = extract_document(_quote_doc(), "quote.docx")
        supplier = [f for f in doc.fields if f.field_name == "supplier_name"]
        assert supplier, "real supplier name was not extracted"
        assert supplier[0].normalized_value == "Al Noor Trading Est."

    def test_unrecognised_lines_are_preserved(self):
        doc = extract_document(_quote_doc(), "quote.docx")
        # The heading and the free-form remark are kept rather than dropped.
        assert len(doc.unclassified_lines) >= 1

    def test_identifiers_do_not_carry_amounts(self):
        doc = extract_document(_quote_doc(), "quote.docx")
        by_name = {f.field_name: f for f in doc.fields}
        assert by_name["quote_number"].is_amount is None
        assert by_name["supplier_name"].is_amount is None


# ── PDF ─────────────────────────────────────────────────────────────────────

class TestPdf:
    @pytest.fixture(scope="class")
    def invoice_pdf(self) -> bytes:
        pytest.importorskip("reportlab")
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        c.drawString(60, 780, "TAX INVOICE")
        c.drawString(60, 758, "Invoice No: INV-2026-0007")
        c.drawString(60, 736, "Supplier: Al Noor Trading Est.")
        c.drawString(60, 714, "Invoice Date: 2026-09-30")
        c.drawString(60, 692, "Total: 1,250.00 SAR")
        c.showPage()
        c.save()
        return buf.getvalue()

    def test_extracts_invoice_fields(self, invoice_pdf):
        doc = extract_document(invoice_pdf, "invoice.pdf")
        assert doc.pages == 1
        names = {f.field_name for f in doc.fields}
        assert "invoice_number" in names
        assert "supplier_name" in names

    def test_total_amount_parsed(self, invoice_pdf):
        doc = extract_document(invoice_pdf, "invoice.pdf")
        assert float(doc.amount_field("total")) == 1250.00

    def test_tax_invoice_title_is_not_a_tax_amount(self, invoice_pdf):
        doc = extract_document(invoice_pdf, "invoice.pdf")
        tax = [f for f in doc.fields if f.field_name == "tax"]
        assert not tax, f"title line misread as a tax field: {[f.normalized_value for f in tax]}"

    def test_pdf_sniffed_by_magic_bytes(self, invoice_pdf):
        # Extension is deliberately wrong; content should still route to PDF.
        doc = extract_document(invoice_pdf, "invoice.dat")
        assert doc.doc_type == "pdf"


class TestHonestFailure:
    def test_scanned_pdf_is_not_guessed(self):
        """No text layer => reported empty with a limitation, never invented."""
        blank = _minimal_pdf("")
        doc = extract_document(blank, "scanned.pdf")
        assert doc.pages == 1
        assert doc.is_empty is True
        assert any("OCR" in lim for lim in doc.limitations)
        assert doc.fields == []

    def test_unsupported_image_is_categorised(self):
        with pytest.raises(IngestionError) as exc:
            extract_document(b"\x89PNG\r\n\x1a\n", "scan.png")
        assert exc.value.category is ErrorCategory.UNSUPPORTED_FORMAT

    def test_malformed_pdf_is_parse_error(self):
        with pytest.raises(IngestionError) as exc:
            extract_document(b"%PDF-1.4 but truncated garbage", "bad.pdf")
        assert exc.value.category is ErrorCategory.PARSE_ERROR

    def test_extraction_method_is_deterministic(self):
        doc = extract_document(_minimal_pdf("Invoice No: INV-9"), "x.pdf")
        assert all(f.method is ExtractionMethod.DOCUMENT_TEXT for f in doc.fields)


# ── Plain text ──────────────────────────────────────────────────────────────

class TestPlainText:
    def test_reads_labelled_fields(self):
        doc = extract_document(
            b"Invoice No: INV-T-1\nTotal: 99.50 SAR\nrandom prose line\n", "note.txt"
        )
        assert {f.field_name for f in doc.fields} >= {"invoice_number", "total"}
        assert float(doc.amount_field("total")) == 99.50

    def test_empty_text_is_reported(self):
        doc = extract_document(b"   \n  ", "blank.txt")
        assert doc.is_empty is True
        assert doc.limitations